import re
import os
import inspect
from openpyxl import load_workbook

from . import log
from . import *

def write_excel_report(olde_sheet_url, report_file_path=None, report_dir_path=None):

  # -----------------------------------------------------------------------------
  # Modifify SimulateTakeOff
  # -----------------------------------------------------------------------------

  log.info(f'Modifying SimulateTakeOff (removing asserts)...')

  # Get source code
  code = inspect.getsource(SimulateTakeOff)

  # Get rid of asserts
  modified_code = []

  lines = iter(code.splitlines())
  for line in lines:
    # Change the class name
    if re.match(r'^class SimulateTakeOff', line):
      line = re.sub(r'class SimulateTakeOff', r'class ModifiedSimulateTakeOff', line)

    # Get rid of asserts
    if re.match(r'^\s*assert', line):
      line = re.sub(r'(^\s*)', r'\1if False: ', line)
    modified_code.append(line)

  modified_code = '\n'.join(modified_code)

  # Optionally save the modified code
  if True:
    with open(os.path.join(get_option('report_dir'), 'ModifiedSimulateTakeOff.py'), 'w', encoding = 'utf-8') as f:
      f.write(modified_code)

  # Load the modified code
  exec(modified_code, globals(), globals())

  # -----------------------------------------------------------------------------
  # Load Ye Olde Sheet
  # -----------------------------------------------------------------------------
  log.info(f'Retrieving the Olde Sheet...')
  olde_workbook = load_workbook(io.BytesIO(get_workbook(olde_sheet_url)), data_only = True)
  olde_sheet = olde_workbook['CES production function']

  t_step = olde_sheet['A8'].value
  n_labour_tasks = int(olde_sheet['JT8'].value)

  for c in reversed(olde_sheet['A']):
    if c.value is not None:
      last_year = c.value
      last_row = c.row
      break

  # -----------------------------------------------------------------------------
  # Simulation
  # -----------------------------------------------------------------------------

  # Retrieve parameter estimates
  log.info(f'Retrieving parameters...')
  parameter_table = get_parameter_table()
  parameters = {parameter : row['Best guess'] for parameter, row in parameter_table.iterrows()}

  parameters['n_labour_tasks'] = n_labour_tasks
  parameters['t_end'] = last_year + t_step

  model = ModifiedSimulateTakeOff(**parameters, t_step=t_step)

  log.info(f'Running the simulation...')
  model.run_simulation()

  # -----------------------------------------------------------------------------
  # Compare the results and write the report
  # -----------------------------------------------------------------------------
  def load_sheet_range(range):
    to_float = lambda x: np.nan if (x is None or x == '#REF!') else float(x)

    if re.match(r'^[A-Z]+$', range):
      col = range
      range = f'{col}54:{col}{last_row}'

    if ':' in range:
      # This is a real range
      array = []
      for row in olde_sheet[range]:
        for cell in row:
          array.append(to_float(cell.value))
      return np.array(array)
    else:
      # This is just a single cell
      return to_float(olde_sheet[range].value)

  clip_idx = model.t_idx
  clip_year = get_option('t_end')
  if clip_year is not None:
    clip_idx = min(clip_idx, model.time_to_index(clip_year))

  def clip(array):
    return array[:clip_idx]

  def yearly_growth_rate(array):
    steps_per_year = round(1/t_step)
    return np.log(array[steps_per_year:]/array[:-steps_per_year])

  olde = {
    'agi_year'                   : load_sheet_range('BA54'),
    'full_rnd_automation_year'   : load_sheet_range('BB54'),
    'fast_growth_year'           : load_sheet_range('BC54'),
    'wakeup_year'                : float(model.index_to_time(np.argmax(load_sheet_range('B')))),

    'timesteps'                  : clip(load_sheet_range('A')),
    'gwp'                        : clip(load_sheet_range('P')),
    'frac_gwp_compute'           : clip(load_sheet_range('D')),
    'frac_compute_training'      : clip(load_sheet_range('BR')),
    'compute'                    : clip(10**load_sheet_range('G')),
    'software'                   : clip(load_sheet_range('F')),
    'biggest_training_run'       : clip(10**load_sheet_range('BT')),
    'hardware_performance'       : clip(10**load_sheet_range(f'AD53:AD{last_row-1}')),

    'frac_tasks_automated_goods' : clip((load_sheet_range('JQ') - load_sheet_range('JQ52'))/load_sheet_range('JT8')) * 100,
    'frac_tasks_automated_rnd'   : clip((load_sheet_range('QH') - load_sheet_range('JQ52'))/load_sheet_range('JT8')) * 100,
  }

  olde['compute_investment'] = olde['gwp'] * olde['frac_gwp_compute'] * t_step

  modern = {
    'agi_year'                   : float(model.index_to_time(np.argmax(clip(model.frac_tasks_automated_goods)))),
    'full_rnd_automation_year'   : float(model.index_to_time(np.argmax(clip(model.frac_tasks_automated_rnd)))),
    'fast_growth_year'           : float(model.index_to_time(np.argmax(yearly_growth_rate(clip(model.gwp)) > 0.20))),
    'wakeup_year'                : model.rampup_start,

    'timesteps'                  : clip(model.timesteps),
    'gwp'                        : clip(model.gwp),
    'frac_gwp_compute'           : clip(model.frac_gwp_compute),
    'frac_compute_training'      : clip(model.frac_compute_training),
    'compute'                    : clip(model.compute),
    'software'                   : clip(model.software),
    'biggest_training_run'       : clip(model.biggest_training_run),
    'hardware_performance'       : clip(model.hardware_performance),

    'frac_tasks_automated_goods' : clip(model.frac_tasks_automated_goods) * 100,
    'frac_tasks_automated_rnd'   : clip(model.frac_tasks_automated_rnd) * 100,
  }

  modern['compute_investment'] = modern['gwp'] * modern['frac_gwp_compute'] * t_step

  for key in ['software', 'hardware_performance', 'gwp']:
    olde[f'{key}_growth'] = yearly_growth_rate(olde[key]) * 100
    modern[f'{key}_growth'] = yearly_growth_rate(modern[key]) * 100

  year_variables = {
    'AGI': 'agi_year',
    'Wake-up': 'wakeup_year',
    'Full R&D automation': 'full_rnd_automation_year',
    'First 20% growth': 'fast_growth_year',
  }

  state_variables = {
    'GWP': 'gwp',
    'Compute': 'compute',
    'Frac compute training': 'frac_compute_training',
    'Biggest training run': 'biggest_training_run',
    'Hardware performance': 'hardware_performance',
    'Compute investment': 'compute_investment',

    'Frac tasks automated goods': 'frac_tasks_automated_goods',
    'Frac tasks automated R&D' : 'frac_tasks_automated_rnd',

    'Yearly software growth (%)': 'software_growth',
    'Yearly hardware performance growth (%)': 'hardware_performance_growth',
    'Yearly GWP growth (%)': 'gwp_growth',
  }

  # Write report
  log.info(f'Writing the report...')
  if report_file_path is None:
    report_file_path = 'excel_comparison.html'

  report = Report(report_file_path=report_file_path, report_dir_path=report_dir_path)
  report.make_tables_scrollable = False

  # Plot compute decomposition
  plot_compute_decomposition_comparison(0, clip_idx, model, olde, modern, normalize = True)
  report.add_figure()

  # Add tables
  year_variables_table = {label: [olde[var], modern[var]] for label, var in year_variables.items()}
  year_variables_table = pd.DataFrame(year_variables_table, index = ['Excel', 'Python'])
  report.add_data_frame(year_variables_table)

  for label, var in state_variables.items():
    container = report.add_html('<div class="header-container"></div>')
    header = report.add_header(label, level = 4, parent = container)
    table = {f'{year:.1f}': [olde[var][i], modern[var][i]] for i, year in enumerate(clip(model.timesteps)[:len(modern[var])])}
    table = pd.DataFrame(table, index = ['Excel', 'Python'])
    html_table = report.add_data_frame(table)

    # Make all columns the same width in the most horribly hacky way
    col_width = 100
    html_table.find('.//table').attrib['style'] = f'table-layout: fixed; width: {col_width * len(modern[var])}px'
    for th in html_table.findall('.//th'):
      th.attrib['style'] = f'width: {col_width}px'

  # Add inputs
  report.add_header("Inputs", level = 3)
  report.add_paragraph(f"Olde sheet: <a href='{olde_sheet_url}'>{olde_sheet_url}</a>")
  report.add_paragraph(f"Python: <a href='{get_option('param_table_url')}'>{get_option('param_table_url')}</a>")
  report.add_data_frame(pd.DataFrame(parameters, index = ['Best guess']).transpose())

  report.add_html('''
    <style>
      .header-container {
        width: 100%;
      }

      h4 {
        position: sticky;
        left: 1em;
        display: inline-block;
        margin-bottom: 2px;
      }
    </style>
  ''', parent = report.head)

  report_path = report.write()
  log.info(f'Report stored in {report_path}')

  return model

def plot_compute_decomposition(
    start_idx, end_idx, reference_idx,
    timesteps, compute_investment, hardware_performance, software, frac_compute_training,
    linestyle = '-', ylim = [None, None], normalize = False):

  clip = lambda x: x[start_idx:end_idx]
  clip_and_normalize = lambda x: clip(x)/x[reference_idx]

  transform = clip_and_normalize if normalize else clip

  plt.plot(clip(timesteps), transform(compute_investment), label = '$ on FLOP globally', color = 'blue', linestyle = linestyle)
  plt.plot(clip(timesteps), transform(hardware_performance), label = 'Hardware quality', color = 'orange', linestyle = linestyle)
  plt.plot(clip(timesteps), transform(software), label = 'Software', color = 'green', linestyle = linestyle)
  plt.plot(clip(timesteps), transform(frac_compute_training), label = 'Fraction global FLOP on training', color = 'red', linestyle = linestyle)
  plt.yscale('log')

  if ylim[0] is not None: plt.ylim(bottom = ylim[0])
  if ylim[1] is not None: plt.ylim(top = ylim[1])

def plot_compute_decomposition_comparison(start_idx, end_idx, model, olde, modern, normalize = False):
  plt.figure(figsize=(14, 8), dpi=80)

  #reference_idx = model.time_to_index(model.rampup_start) if model.rampup_start is not None else start_idx
  reference_idx = 0

  ylim = [1e-1, 1e4] if normalize else [None, 1e21]

  plot_compute_decomposition(
      start_idx, end_idx, reference_idx,
      modern['timesteps'], modern['compute_investment'], modern['hardware_performance'], modern['software'], modern['frac_compute_training'],
      linestyle = '-', ylim = ylim, normalize = normalize,
  )

  model._plot_vlines()

  plot_compute_decomposition(
      start_idx, end_idx, reference_idx,
      olde['timesteps'], olde['compute_investment'], olde['hardware_performance'], olde['software'], olde['frac_compute_training'],
      linestyle = '--', ylim = ylim, normalize = normalize,
  )

  handles, labels = plt.gca().get_legend_handles_labels() # Hacky legend

  draw_oom_lines()

  # Plot the legend

  vline_count = 0
  if model.rampup_start: vline_count += 1
  if model.rampup_mid:   vline_count += 1
  if model.agi_year:     vline_count += 1

  new_handles = handles[:4]
  new_labels = labels[:4]
  vline_handles = handles[4:4+vline_count]
  vline_labels = labels[4:4+vline_count]
  old_handles = handles[4+vline_count:]
  old_labels = labels[4+vline_count:]

  old_legend = plt.legend(old_handles, old_labels, bbox_to_anchor = (1.02, 1), borderaxespad = 0, title = 'Olde model')
  new_legend = plt.legend(new_handles, new_labels, bbox_to_anchor = (1.02, 0.80), borderaxespad = 0, title = 'New model')
  vline_legend = plt.legend(vline_handles, vline_labels, bbox_to_anchor = (1.02, 0.60), borderaxespad = 0)
  plt.legend().remove()

  old_legend.set_frame_on(False)
  new_legend.set_frame_on(False)
  vline_legend.set_frame_on(False)

  old_legend._legend_box.align = "left"
  new_legend._legend_box.align = "left"
  vline_legend._legend_box.align = "left"

  plt.gcf().add_artist(old_legend)
  plt.gcf().add_artist(new_legend)
  plt.gcf().add_artist(vline_legend)


if __name__ == '__main__':
  parser = init_cli_arguments()

  parser.add_argument(
    "-s",
    "--olde-sheet",
    type=str,
    default='https://docs.google.com/spreadsheets/d/1rMMkUsuxzuRqa9R5Foqu7XSinvU6-zow2FzX96PmSrI/edit?usp=sharing',
  )

  args = handle_cli_arguments(parser)

  write_excel_report(
    report_file_path=args.output_file,
    report_dir_path=args.output_dir,
    olde_sheet_url=args.olde_sheet,
  )
